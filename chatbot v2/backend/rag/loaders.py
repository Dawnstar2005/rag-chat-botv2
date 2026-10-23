"""Turn uploaded files into plain text, page by page."""

import csv
import io
import json
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path


class DocumentLoadError(Exception):
    """The file could not be read; the message is shown to the user."""


@dataclass
class Page:
    text: str
    number: int | None = None  # 1-based page number for PDFs, None otherwise


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".json", ".html", ".htm"}


def decode_text(data: bytes) -> str:
    """Decode bytes, trying UTF-8 first, then the Cyrillic Windows code page."""
    for encoding in ("utf-8-sig", "utf-16", "cp1251"):
        try:
            text = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        # utf-16 "succeeds" on almost anything; only accept it with a BOM.
        if encoding == "utf-16" and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
            continue
        return text
    return data.decode("latin-1")


def _load_pdf(data: bytes) -> list[Page]:
    from pypdf import PasswordType, PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        # decrypt() reports a wrong password by its return value, not by raising.
        if reader.is_encrypted and reader.decrypt("") == PasswordType.NOT_DECRYPTED:
            raise DocumentLoadError("PDF файл нууц үгээр хамгаалагдсан байна.")
        pages = []
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                pages.append(Page(text=text, number=index))
        return pages
    except DocumentLoadError:
        raise
    except Exception as exc:  # malformed files raise far more than PdfReadError (TypeError, ValueError...)
        raise DocumentLoadError(f"PDF файлыг уншиж чадсангүй: {exc}") from exc


def _load_docx(data: bytes) -> list[Page]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentLoadError("DOCX файлыг уншиж чадсангүй.") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells, previous = [], None
            for cell in row.cells:
                if cell is not previous:  # a merged cell is listed once per grid column it spans
                    cells.append(cell.text.strip())
                previous = cell
            if any(cells):
                parts.append(" | ".join(cells))
    return [Page(text="\n\n".join(parts))]


def _load_csv(data: bytes) -> list[Page]:
    text = decode_text(data)
    try:
        rows = list(csv.reader(io.StringIO(text)))
    except csv.Error:  # e.g. a field over the csv module's 128 KB limit
        return [Page(text=text)]
    if not rows:
        return []
    header, body = rows[0], rows[1:]
    lines = []
    for row in body:
        pairs = [f"{h}: {v}" for h, v in zip(header, row) if v.strip()]
        lines.append("; ".join(pairs))
    return [Page(text="\n".join(lines) if lines else ", ".join(header))]


def _load_json(data: bytes) -> list[Page]:
    text = decode_text(data)
    try:
        parsed = json.loads(text)
        text = json.dumps(parsed, ensure_ascii=False, indent=2)
    except (ValueError, RecursionError):  # invalid, or nested too deeply to parse: keep the raw text
        pass
    return [Page(text=text)]


class _HTMLText(HTMLParser):
    _SKIP = {"script", "style", "noscript", "head"}
    _BLOCK = {
        "p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article",
        "ul", "ol", "dl", "dt", "dd", "table", "blockquote", "pre", "hr", "header", "footer",
        "nav", "main", "aside", "figure", "figcaption",
    }
    _CELL = {"td", "th"}

    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip_depth = 0
        self._row_has_cell = False

    def handle_starttag(self, tag, attrs):
        if tag in self._SKIP:
            self._skip_depth += 1
        elif tag == "body":
            self._skip_depth = 0  # <head> may be left unclosed
        elif tag in self._CELL:
            # Without a separator "<td>A</td><td>B</td>" would become the single word "AB".
            self.parts.append(" | " if self._row_has_cell else " ")
            self._row_has_cell = True
        elif tag in self._BLOCK:
            if tag == "tr":
                self._row_has_cell = False
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in self._SKIP and self._skip_depth:
            self._skip_depth -= 1
        elif tag in self._BLOCK:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self._skip_depth:
            self.parts.append(data)


def _load_html(data: bytes) -> list[Page]:
    parser = _HTMLText()
    parser.feed(decode_text(data))
    parser.close()  # flushes trailing text the parser holds back (e.g. ending in "R&D")
    return [Page(text="".join(parser.parts))]


def load_document(filename: str, data: bytes) -> list[Page]:
    """Extract text pages from a file. Raises DocumentLoadError on failure."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        allowed = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentLoadError(f"'{ext or filename}' төрлийн файл дэмжигдэхгүй. Дэмжих төрөл: {allowed}")

    if ext == ".pdf":
        pages = _load_pdf(data)
    elif ext == ".docx":
        pages = _load_docx(data)
    elif ext == ".csv":
        pages = _load_csv(data)
    elif ext == ".json":
        pages = _load_json(data)
    elif ext in {".html", ".htm"}:
        pages = _load_html(data)
    else:
        pages = [Page(text=decode_text(data))]

    pages = [p for p in pages if p.text.strip()]
    if not pages:
        hint = " (скан хийсэн зураг PDF байж магадгүй)" if ext == ".pdf" else ""
        raise DocumentLoadError(f"Файлаас текст олдсонгүй{hint}.")
    return pages
